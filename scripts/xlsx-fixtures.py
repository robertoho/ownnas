"""Writes sample workbooks for scripts/test-xlsx-roundtrip.mjs into the folder given.

Needs openpyxl (and Pillow, for the picture). Each workbook has what spreadsheet apps put in
files and the editor must keep: styles, merged cells, frozen panes, a table, conditional
formatting, data validation, comments, links, a picture, a chart, names, and sheets that refer
to each other.
"""

import io
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


def rich(path):
    book = Workbook()
    sales = book.active
    sales.title = "Sales"
    sales.append(["Item", "Qty", "Price", "Total"])
    for item, qty, price in [("Pen", 3, 1.5), ("Ink", 5, 4.0), ("Pad", 2, 2.25), ("Cap", 7, 0.5)]:
        row = sales.max_row + 1
        sales.append([item, qty, price, f"=B{row}*C{row}"])
    sales["F1"] = "Grand total"
    sales["G1"] = "=SUM(D2:D5)"
    sales["A1"].font = Font(bold=True, color="FFFFFF")
    sales["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin")
    sales["G1"].border = Border(top=thin, bottom=thin, left=thin, right=thin)
    sales["C2"].number_format = '#,##0.00 "€"'
    table = Table(displayName="Stock", ref="A1:D5")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sales.add_table(table)
    sales.freeze_panes = "A2"
    sales.conditional_formatting.add("B2:B5", CellIsRule(operator="greaterThan", formula=["4"], fill=PatternFill("solid", bgColor="FFC7CE")))
    sales.conditional_formatting.add("C2:C5", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="63BE7B"))
    sales.conditional_formatting.add("A2:A5", FormulaRule(formula=['$B2>5'], font=Font(bold=True)))
    choice = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True, showInputMessage=True, prompt="Pick one")
    sales.add_data_validation(choice)
    choice.add("E2:E5")
    sales["A3"].comment = Comment("Running low", "Ana")
    sales["F3"] = "Site"
    sales["F3"].hyperlink = "https://example.com/"
    sales.merge_cells("F5:G6")
    sales["F5"] = "Merged note"
    chart = BarChart()
    chart.title = "Totals"
    chart.add_data(Reference(sales, min_col=4, min_row=1, max_row=5), titles_from_data=True)
    sales.add_chart(chart, "I2")
    try:
        from PIL import Image as Picture
        data = io.BytesIO()
        Picture.new("RGB", (40, 20), (200, 30, 30)).save(data, format="PNG")
        data.seek(0)
        sales.add_image(Image(data), "I20")
    except ImportError:
        pass

    summary = book.create_sheet("Summary")
    summary["A1"] = "Total"
    summary["B1"] = "=Sales!G1"
    summary["A2"] = "Items"
    summary["B2"] = "=COUNTA(Sales!A2:A5)"
    book.defined_names["Grand"] = DefinedName("Grand", attr_text="Sales!$G$1")
    book.save(path)


def main():
    folder = Path(sys.argv[1])
    folder.mkdir(parents=True, exist_ok=True)
    rich(folder / "rich.xlsx")


if __name__ == "__main__":
    main()
